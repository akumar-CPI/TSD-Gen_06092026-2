#!/usr/bin/env python3
"""
generate_tsd.py

Reads the JSON produced by parse_iflow.py and renders a Technical
Specification Document (.docx) that follows the FICOPROC TSD template's
table-based layout.

Usage:
    python generate_tsd.py <parsed.json> <output.docx>
"""
import sys
import json
import re
import xml.etree.ElementTree as ET
from collections import OrderedDict
import logging

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from schema_map import PALLET_SCHEMAS, ADAPTER_SCHEMAS

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

FONT_NAME = "Ubuntu"
HEADER_FILL = "1CADE4"
TITLE_COLOR = RGBColor(0x1C, 0xAD, 0xE4)

NESTED_TABLE_LABELS = {
    "headerTable": "Message Header - Headers",
    "propertyTable": "Exchange Property - Properties",
}
PREFERRED_COLUMNS = ["Action", "Name", "Type", "Datatype", "Value", "Default"]


# --------------------------------------------------------------------------
# Low level docx helpers
# --------------------------------------------------------------------------

def set_run_font(run, size, bold=False, color=None):
    run.font.name = FONT_NAME
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    # set fonts for ASCII, hAnsi and eastAsia so Word on all platforms picks it up
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_heading(doc, text, level):
    sizes = {1: 18, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, sizes.get(level, 11), bold=True,
                 color=TITLE_COLOR if level <= 2 else RGBColor(0, 0, 0))
    return p


def add_band_row(table, text, span=2):
    row = table.add_row()
    cell = row.cells[0]
    for i in range(1, span):
        cell = cell.merge(row.cells[i])
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    shade_cell(cell, HEADER_FILL)


def set_multiline_text(paragraph, text):
    lines = (text or "").split("\n")
    run = paragraph.add_run(lines[0])
    set_run_font(run, 9, bold=False)
    for line in lines[1:]:
        run.add_break()
        run2 = paragraph.add_run(line)
        set_run_font(run2, 9, bold=False)
    return run


def add_label_value_row(table, label, value):
    row = table.add_row()
    run = row.cells[0].paragraphs[0].add_run(label)
    set_run_font(run, 10, bold=True)
    set_multiline_text(row.cells[1].paragraphs[0], value)


def new_table(doc):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    # try to set column widths (python-docx can be inconsistent here)
    try:
        table.columns[0].width = Cm(5.5)
        table.columns[1].width = Cm(10.5)
    except Exception:
        pass
    # CRITICAL: Word (and python-docx) merges two <w:tbl> elements that
    # sit back-to-back with nothing between them into ONE visual table.
    # Every table this generator creates must be followed by a paragraph
    # before anything else is added, or repeated pallet-function tables
    # will visually collapse into a single giant table. This spacer is
    # added immediately so callers can never forget it.
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_multi_column_table(doc, rows, columns=None, container_cell=None):
    """Create a multi-column table. If container_cell is provided, the
    table is added inside that cell (nested table). Otherwise it's added
    at document level.
    """
    if columns is None:
        columns = []
        for r in rows:
            for k in r.keys():
                if k not in columns:
                    columns.append(k)
        ordered = [c for c in PREFERRED_COLUMNS if c in columns]
        ordered += [c for c in columns if c not in ordered]
        columns = ordered or ["Value"]

    if container_cell is not None:
        table = container_cell.add_table(rows=0, cols=len(columns))
    else:
        table = doc.add_table(rows=0, cols=len(columns))
        # when adding at document level, keep the usual style/spacer behavior
        table.style = "Table Grid"

    # apply style for both nested and top-level tables
    try:
        table.style = table.style or "Table Grid"
    except Exception:
        pass

    hdr = table.add_row()
    for i, col in enumerate(columns):
        run = hdr.cells[i].paragraphs[0].add_run(col)
        set_run_font(run, 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(hdr.cells[i], HEADER_FILL)
    for r in rows:
        row = table.add_row()
        for i, col in enumerate(columns):
            set_multiline_text(row.cells[i].paragraphs[0], r.get(col, ""))

    if container_cell is None:
        # only add the document-level spacer when table is top-level
        doc.add_paragraph().paragraph_format.space_after = Pt(6)  # see new_table() note
    return table


# --------------------------------------------------------------------------
# Nested "table-inside-a-property" detection & parsing (CPI-specific)
# --------------------------------------------------------------------------

_ROW_HINT_RE = re.compile(r"<\s*row\b", re.IGNORECASE)


def looks_like_nested_table(value):
    return bool(value) and bool(_ROW_HINT_RE.search(value))


def parse_nested_table(value):
    wrapped = f"<root>{value}</root>"
    try:
        root = ET.fromstring(wrapped)
    except ET.ParseError:
        return None
    rows = []
    for row_el in root.findall(".//row"):
        row_data = OrderedDict()
        for idx, cell in enumerate(row_el.findall("cell")):
            cid = cell.get("id") or cell.get("name") or cell.get("key")
            if not cid:
                cid = PREFERRED_COLUMNS[idx] if idx < len(PREFERRED_COLUMNS) else f"Col{idx+1}"
            row_data[cid] = cell.text or ""
        if row_data:
            rows.append(row_data)
    return rows or None


def extract_nested_tables(properties):
    """Pull out any CPI 'table inside a property' values (headerTable,
    propertyTable, ...) from a raw properties dict, returning
    {key: [row_dicts]} for ones that parsed successfully."""
    nested = OrderedDict()
    for k, v in properties.items():
        if looks_like_nested_table(v):
            rows = parse_nested_table(v)
            if rows:
                nested[k] = rows
    return nested


# --------------------------------------------------------------------------
# Field resolution (template-only, no fallback dumping)
# --------------------------------------------------------------------------

def render_bool(value):
    truthy = (value or "").strip().lower() in ("true", "yes", "1")
    return "\u2611" if truthy else "\u2610"


def render_checkbox_value(options, current_value):
    cur = (current_value or "").strip().lower()
    if isinstance(options, dict):
        parts = []
        for raw_value, label in options.items():
            matched = cur == str(raw_value).strip().lower()
            box = "\u2611" if matched else "\u2610"
            parts.append(f"{box} {label}")
        return "   ".join(parts)
    parts = []
    for opt in options:
        matched = cur == opt.strip().lower() or (cur and opt.strip().lower() in cur)
        box = "\u2611" if matched else "\u2610"
        parts.append(f"{box} {opt}")
    return "   ".join(parts)


def resolve_field(context, field_spec):
    if field_spec == "__NAME__":
        return context.get("__name__", "")
    if isinstance(field_spec, dict):
        if "bool" in field_spec:
            return render_bool(context.get(field_spec["bool"], ""))
        prop = field_spec.get("prop")
        options = field_spec.get("checkbox", [])
        return render_checkbox_value(options, context.get(prop, ""))
    if isinstance(field_spec, (list, tuple)):
        # Ordered list of alternative property keys (some CPI adapter
        # versions store the same concept under a different key, e.g.
        # "Address" vs "httpAddressWithoutQuery") - first non-empty wins.
        for key in field_spec:
            val = context.get(key)
            if val:
                return val
        return ""
    return context.get(field_spec, "")


# --- Helpers to find body text and external params ---------------------------------
EXTERNAL_PARAM_KEYS = [
    "externalParameters", "externalParameterTable", "externalParametersTable",
    "externalParametersList", "adapterSpecificParameters", "externalParams",
    "externalParameter"
]
BODY_FALLBACK_KEYS = [
    "bodyContent", "body", "messageBody", "content", "MessageBody",
    "bodyText", "soapBody", "payload"
]


def extract_text_from_xml(maybe_xml):
    """If value contains XML, strip tags and return concatenated text nodes; else return original string."""
    if not maybe_xml or "<" not in maybe_xml:
        return maybe_xml or ""
    wrapped = f"<root>{maybe_xml}</root>"
    try:
        root = ET.fromstring(wrapped)
        return "".join(root.itertext()).strip()
    except ET.ParseError:
        # not well-formed XML - return original
        return maybe_xml


def find_body_in_context(context):
    for key in BODY_FALLBACK_KEYS:
        v = context.get(key)
        if v:
            txt = extract_text_from_xml(v)
            logger.debug(f"Found body using key '{key}' (len={len(txt)})")
            return txt
    # as a last resort, check for single large property that may contain the body
    for k, v in context.items():
        if isinstance(v, str) and len(v) > 200 and "<row" not in v:
            txt = extract_text_from_xml(v)
            logger.debug(f"Heuristic: using key '{k}' as body (len={len(txt)})")
            return txt
    return ""


# --------------------------------------------------------------------------
# Main per-element renderer
# --------------------------------------------------------------------------

def render_schema_table(doc, schema, name, properties, elem_id="", routes=None):
    """Render one instance's table: only the fields the schema defines,
    plus any nested CPI table properties relevant to this element, plus
    (for Router) its route conditions pulled from sequenceFlow data."""
    context = dict(properties)
    # Some real exports leave certain steps (Router/ExclusiveGateway in
    # particular) with no name set - fall back to the element's own id so
    # multiple instances of the same type are still distinguishable.
    context["__name__"] = name or elem_id or "(unnamed)"

    # debug: show available property keys when verbose
    logger.debug(f"Rendering element {context['__name__']} ({elem_id}) properties: {list(properties.keys())}")

    table = new_table(doc)
    for section_name, fields in schema["sections"]:
        add_band_row(table, section_name)
        for label, field_spec in fields:
            # resolve value normally, then apply ContentModifier body fallback
            val = resolve_field(context, field_spec)
            # if this field is the Body field commonly named 'bodyContent',
            # and it's empty, try a list of fallback property keys
            is_body_field = False
            if field_spec == "bodyContent":
                is_body_field = True
            elif isinstance(field_spec, (list, tuple)) and "bodyContent" in field_spec:
                is_body_field = True

            if is_body_field and not val:
                val = find_body_in_context(context)

            add_label_value_row(table, label, val)
            # log missing expected fields for easier debugging
            if not val and field_spec not in ("__NAME__",):
                logger.debug(f"Field '{field_spec}' for label '{label}' is empty for element {context['__name__']} ({elem_id})")

    # Insert nested CPI tables inline in the value cell of a new row so they
    # visually belong to the same element block (e.g. Content Modifier).
    nested = extract_nested_tables(properties)
    for key, rows in nested.items():
        label = NESTED_TABLE_LABELS.get(key, key)
        # add a label/value row and then put the nested multi-column table
        # inside the value cell of that newly added row
        add_label_value_row(table, label, "")
        container_cell = table.rows[-1].cells[1]
        add_multi_column_table(doc, rows, container_cell=container_cell)

    # also try to render External Parameters (common property names)
    for ep_key in EXTERNAL_PARAM_KEYS:
        if ep_key in properties and properties[ep_key]:
            val = properties[ep_key]
            # if it's a nested CPI table string, parse and render inline
            if looks_like_nested_table(val):
                rows = parse_nested_table(val)
                if rows:
                    add_label_value_row(table, "External Parameters", "")
                    container_cell = table.rows[-1].cells[1]
                    add_multi_column_table(doc, rows, container_cell=container_cell)
                    continue
            # try JSON list/dict
            try:
                parsed = json.loads(val)
                # if it's a list of objects, render column headers as keys
                if isinstance(parsed, list) and parsed:
                    rows = []
                    if isinstance(parsed[0], dict):
                        for item in parsed:
                            row = OrderedDict()
                            for k, v in item.items():
                                row[str(k)] = str(v)
                            rows.append(row)
                        add_label_value_row(table, "External Parameters", "")
                        container_cell = table.rows[-1].cells[1]
                        add_multi_column_table(doc, rows, container_cell=container_cell)
                        continue
                elif isinstance(parsed, dict):
                    rows = [ {k: str(v)} for k, v in parsed.items() ]
                    add_label_value_row(table, "External Parameters", "")
                    container_cell = table.rows[-1].cells[1]
                    add_multi_column_table(doc, rows, container_cell=container_cell)
                    continue
            except Exception:
                # not JSON - just render as text under External Parameters
                add_label_value_row(table, "External Parameters", extract_text_from_xml(str(val)))

    if routes:
        # similarly add the Route Conditions table inline
        add_label_value_row(table, "Route Conditions", "")
        container_cell = table.rows[-1].cells[1]
        add_multi_column_table(doc, routes, columns=["Order", "Route Name", "Conditional Expression", "Default Route"], container_cell=container_cell)


def render_unmapped_table(doc, name, elem_id=""):
    table = new_table(doc)
    add_band_row(table, "General")
    add_label_value_row(table, "Name", name or elem_id or "(unnamed)")


# --------------------------------------------------------------------------
# Section numbering + grouping
# --------------------------------------------------------------------------

def group_in_order(items, key_fn):
    groups = OrderedDict()
    for item in items:
        k = key_fn(item)
        groups.setdefault(k, []).append(item)
    return groups


class Numbering:
    """Produces 'X' and 'X.Y' heading prefixes."""
    def __init__(self):
        self.top = 0
        self.sub = 0

    def next_top(self):
        self.top += 1
        self.sub = 0
        return str(self.top)

    def next_sub(self):
        self.sub += 1
        return f"{self.top}.{self.sub}"


# --------------------------------------------------------------------------
# Main render
# --------------------------------------------------------------------------

def render(data, out_path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.get("package_name", "Integration Flow"))
    set_run_font(run, 24, bold=True, color=TITLE_COLOR)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Technical Specification Document")
    set_run_font(run2, 14, bold=False)
    doc.add_page_break()

    numbering = Numbering()

    # 1. Participants
    n = numbering.next_top()
    add_heading(doc, f"{n}. Participants", level=1)
    table = new_table(doc)
    add_band_row(table, "Participants")
    for participant in data.get("participants", []):
        add_label_value_row(table, participant.get("name", ""),
                             participant.get("properties", {}).get("ifl:type", ""))
    if not data.get("participants"):
        add_label_value_row(table, "(none found)", "")

    # 2. Processes
    n = numbering.next_top()
    add_heading(doc, f"{n}. Process Overview", level=1)
    processes = data.get("processes", [])
    process_groups = group_in_order(processes, lambda p: p["kind"])
    for kind, items in process_groups.items():
        sub = numbering.next_sub()
        add_heading(doc, f"{sub} {kind}", level=2)
        for proc in items:
            table = new_table(doc)
            add_band_row(table, "General")
            add_label_value_row(table, "Name", proc["name"])
            add_label_value_row(table, "ID", proc["id"])

    # 3. Pallet Function Details
    n = numbering.next_top()
    add_heading(doc, f"{n}. Pallet Function Details", level=1)
    pallet_elements = data.get("pallet_elements", [])
    PLUMBING_ONLY_KEYS = {"activityType", "cmdVariantUri", "componentVersion"}

    def is_bare_boundary_event(e):
        # A plain start/end event carries nothing beyond activityType/
        # cmdVariantUri/componentVersion. A Timer start event (which also
        # uses activityType "StartEvent" in some exports) carries real
        # scheduler properties too, so it's kept rather than excluded.
        if e["xml_tag"] not in ("startEvent", "endEvent"):
            return False
        if not e["activity_type"]:
            return True
        return set(e["properties"].keys()) <= PLUMBING_ONLY_KEYS

    pallet_elements = [e for e in pallet_elements if not is_bare_boundary_event(e)]
    groups = group_in_order(pallet_elements, lambda e: e.get("resolved_type") or e["activity_type"] or e["xml_tag"])
    if not groups:
        doc.add_paragraph("No pallet function elements were found in this iFlow.")
    for resolved_type, elements in groups.items():
        schema = PALLET_SCHEMAS.get(resolved_type)
        title = schema["title"] if schema else resolved_type
        sub = numbering.next_sub()
        add_heading(doc, f"{sub} {title}", level=2)
        for elem in elements:
            if schema:
                render_schema_table(doc, schema, elem["name"], elem["properties"],
                                     elem_id=elem.get("id", ""), routes=elem.get("routes"))
            else:
                render_unmapped_table(doc, elem["name"], elem_id=elem.get("id", ""))

    # 4. Connectivity (Adapters)
    n = numbering.next_top()
    add_heading(doc, f"{n}. Connectivity", level=1)
    adapters = data.get("adapters", [])
    if not adapters:
        doc.add_paragraph("No sender/receiver channels were found in this iFlow.")
    a_groups = group_in_order(
        adapters, lambda a: f"{(a['component_type'] or '').lower()}|{(a['direction'] or '').lower()}"
    )
    for key, items in a_groups.items():
        schema = ADAPTER_SCHEMAS.get(key)
        first = items[0]
        title = schema["title"] if schema else f"{first['component_type'] or 'Adapter'} ({first['direction'] or 'n/a'})"
        sub = numbering.next_sub()
        add_heading(doc, f"{sub} {title}", level=2)
        for ad in items:
            if schema:
                render_schema_table(doc, schema, ad["name"], ad["properties"], elem_id=ad.get("id", ""))
            else:
                render_unmapped_table(doc, ad["name"], elem_id=ad.get("id", ""))

    try:
        doc.save(out_path)
        logger.info(f"Wrote {out_path}")
    except Exception:
        logger.exception("Failed to write output docx")
        raise


def main():
    if len(sys.argv) != 3:
        logger.error("Usage: python generate_tsd.py <parsed.json> <output.docx>")
        sys.exit(1)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    render(data, sys.argv[2])


if __name__ == "__main__":
    main()
